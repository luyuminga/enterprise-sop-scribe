#!/usr/bin/env python3
"""生成进场空白表（配套 enterprise-sop-scribe）。

输出到仓库内 templates/excel/：
  - 组织信息填写表.xlsx
  - 个人工作事项填写表.xlsx

依赖：pip install openpyxl
用法：python3 scripts/generate_prefill_forms.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "templates" / "excel"
ORG_NAME = "组织信息填写表.xlsx"
PERSONAL_NAME = "个人工作事项填写表.xlsx"

C_PRIMARY = "1F4E79"
C_HEADER = "2E5C8A"
C_SECTION = "E8EEF4"
C_INPUT = "FFFFFF"  # 待填区：白底，不用淡黄
C_ALT = "F5F8FB"
C_LINE = "A8B8C8"
C_HINT = "666666"

TITLE_FONT = Font(name="微软雅黑", size=15, bold=True, color=C_PRIMARY)
DOC_FONT = Font(name="微软雅黑", size=9, color="888888")
HEADER_FONT = Font(name="微软雅黑", size=9, bold=True, color="FFFFFF")
LABEL_FONT = Font(name="微软雅黑", size=10, bold=True, color="333333")
CELL_FONT = Font(name="微软雅黑", size=10, color="333333")
HINT_FONT = Font(name="微软雅黑", size=9, color=C_HINT)

HEADER_FILL = PatternFill("solid", fgColor=C_HEADER)
SECTION_FILL = PatternFill("solid", fgColor=C_SECTION)
INPUT_FILL = PatternFill("solid", fgColor=C_INPUT)
ALT_FILL = PatternFill("solid", fgColor=C_ALT)

THIN = Border(
    left=Side(style="thin", color=C_LINE),
    right=Side(style="thin", color=C_LINE),
    top=Side(style="thin", color=C_LINE),
    bottom=Side(style="thin", color=C_LINE),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
PAGE_MARGIN = PageMargins(left=0.5, right=0.5, top=0.55, bottom=0.55)

# 填写行默认高度（方便打字）
INPUT_ROW_H = 36
HEADER_ROW_H = 48
LABEL_ROW_H = 30



def move_sheet_to_end(wb, title="填写说明"):
    if title in wb.sheetnames:
        ws = wb[title]
        wb.move_sheet(ws, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(title))


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def paint(ws, r1, r2, c1, c2, *, fill=None, font=None, align=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            if fill is not None:
                cell.fill = fill
            if font is not None:
                cell.font = font
            if align is not None:
                cell.alignment = align
            cell.border = THIN


def merge_write(ws, r, c1, c2, value, *, fill=None, font=None, align=None, height=None):
    paint(ws, r, r, c1, c2, fill=fill, font=font or CELL_FONT, align=align or LEFT)
    cell = ws.cell(r, c1, value)
    if font is not None:
        cell.font = font
    if align is not None:
        cell.alignment = align
    if fill is not None:
        cell.fill = fill
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    if height:
        ws.row_dimensions[r].height = height
    return cell


def setup_sheet(ws, last_col="L", landscape=True, freeze="A4"):
    ws.sheet_view.showGridLines = False
    ws.page_margins = PAGE_MARGIN
    ws.print_title_rows = "1:3"
    if freeze:
        ws.freeze_panes = freeze
    if landscape:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_area = f"A1:{last_col}{ws.max_row}"


def write_header_row(ws, row, headers, *, height=HEADER_ROW_H):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN
    ws.row_dimensions[row].height = height


def blank_input_rows(ws, start, end, ncols, *, height=INPUT_ROW_H, alt=True):
    for r in range(start, end + 1):
        fill = ALT_FILL if (alt and (r - start) % 2) else INPUT_FILL
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c, "")
            cell.fill = fill
            cell.font = CELL_FONT
            cell.alignment = LEFT_TOP
            cell.border = THIN
        ws.row_dimensions[r].height = height


def label_value_row(ws, row, pairs, *, height=LABEL_ROW_H):
    for item in pairs:
        if len(item) == 3:
            text, lc, vc = item
            ve = vc
        else:
            text, lc, vc, ve = item
        cell = ws.cell(row, lc, text)
        cell.fill = SECTION_FILL
        cell.font = LABEL_FONT
        cell.alignment = CENTER
        cell.border = THIN
        for c in range(vc, ve + 1):
            cell = ws.cell(row, c, "")
            cell.fill = INPUT_FILL
            cell.font = CELL_FONT
            cell.alignment = LEFT
            cell.border = THIN
        if ve > vc:
            ws.merge_cells(start_row=row, start_column=vc, end_row=row, end_column=ve)
    ws.row_dimensions[row].height = height


def build_org_workbook():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # ---- 填写说明 ----
    ws = wb.active
    ws.title = "填写说明"
    set_widths(ws, [12, 14, 72, 12, 12])
    merge_write(ws, 1, 1, 5, "组织信息填写表 · 填写说明", fill=SECTION_FILL, font=TITLE_FONT, align=CENTER, height=34)
    merge_write(
        ws, 2, 1, 5,
        "本表由本部门负责人汇总填写；与《个人工作事项填写表》成套使用。白色为待填。",
        font=DOC_FONT, align=CENTER, height=22,
    )
    rows = [
        ("用途", "进场前整理本部门编制、业务量基线、系统与产物，方便后续对接，减少现场空转。"),
        ("谁填", "本部门负责人填写（可指定骨干代为汇总，仍按本部门口径）。"),
        ("成套关系", "负责人填本表；本部门同事各填《个人工作事项填写表》。一份部门表 + 若干个人表，互不替代。"),
        ("脱敏", "勿填账号密码、业务往来方完整姓名/证件号、未脱敏单号；样例可另附脱敏文件。"),
        ("SOP 收集", "本部门同事在各自《个人工作事项填写表》按事项写文档名；由您统一收集文件，打成一个压缩包与本表一并提交。"),
        ("不确定时", "可填区间或写「待确认」；基线允许估算，请在备注注明「估算」。"),
        ("交回时间", "建议进场前 3～5 个工作日交回。"),
        ("本表结构", "①基本信息与配合（含 SOP 收集）  ②岗位编制  ③业务量基线  ④系统与产物清单"),
    ]
    r = 4
    for title, body in rows:
        merge_write(ws, r, 1, 2, title, fill=SECTION_FILL, font=LABEL_FONT, align=CENTER, height=42)
        merge_write(ws, r, 3, 5, body, fill=INPUT_FILL, font=CELL_FONT, align=LEFT_TOP, height=42)
        r += 1
    merge_write(
        ws, r + 1, 1, 5,
        "交回：本部门组织信息 Excel + 本部门 SOP 压缩包（如有）+ 本部门个人表。Excel 建议名：组织信息填写表-公司名-部门名.xlsx；压缩包：标准作业SOP-公司名-部门名.zip",
        font=HINT_FONT, align=LEFT_TOP, height=40,
    )
    setup_sheet(ws, last_col="E", landscape=False, freeze=None)

    # ---- 01基本信息 ----
    ws = wb.create_sheet("01基本信息与配合")
    set_widths(ws, [22, 24, 18, 24, 16, 22, 14, 22])
    merge_write(ws, 1, 1, 8, "组织信息填写表 · 基本信息与配合（部门维度）", fill=SECTION_FILL, font=TITLE_FONT, align=CENTER, height=32)
    merge_write(ws, 2, 1, 8, "由本部门负责人汇总填写；配合同部门《个人工作事项填写表》使用", font=DOC_FONT, align=CENTER, height=20)

    # 预计进场日、调研范围等由项目双方决策人对齐，不放本表
    label_value_row(ws, 3, [("公司名称", 1, 2), ("本部门", 3, 4), ("填表日期", 5, 6), ("部门人数（约）", 7, 8)], height=32)
    label_value_row(ws, 4, [("填写人（负责人）", 1, 2), ("岗位/职务", 3, 4), ("手机/微信", 5, 6), ("邮箱", 7, 8)], height=32)
    label_value_row(ws, 5, [("班制/作息", 1, 2), ("高峰时段", 3, 4)], height=32)

    merge_write(ws, 7, 1, 8, "一、现场配合与约束（本部门可答复）", fill=HEADER_FILL, font=HEADER_FONT, align=LEFT, height=26)
    for text, row in [
        ("进场时可提供哪些材料/演示\n（如脱敏单据、报表样例、系统只读演示）", 8),
        ("方便配合的时间与条件\n（可访谈时段、有无会议室、是否允许旁观操作电脑）", 9),
        ("安全与保密要求\n（可否拍照/录音、资料能否外带）", 10),
        ("近半年系统或流程是否大改\n（是/否 + 简要说明）", 11),
    ]:
        merge_write(ws, row, 1, 2, text, fill=SECTION_FILL, font=LABEL_FONT, align=LEFT_TOP, height=44)
        merge_write(ws, row, 3, 8, "", fill=INPUT_FILL, font=CELL_FONT, align=LEFT_TOP, height=44)

    merge_write(ws, 13, 1, 8, "二、标准作业 SOP 收集（由您统一收集提交）", fill=HEADER_FILL, font=HEADER_FONT, align=LEFT, height=26)
    paint(ws, 14, 14, 1, 2, fill=SECTION_FILL, font=LABEL_FONT, align=LEFT_TOP)
    ws.cell(14, 1, "是否已按个人事项表\n收集齐相关 SOP")
    ws.merge_cells("A14:B14")
    paint(ws, 14, 14, 3, 4, fill=INPUT_FILL, font=CELL_FONT, align=LEFT)
    ws.merge_cells("C14:D14")
    paint(ws, 14, 14, 5, 5, fill=SECTION_FILL, font=LABEL_FONT, align=CENTER)
    ws.cell(14, 5, "压缩包文件名")
    paint(ws, 14, 14, 6, 8, fill=INPUT_FILL, font=CELL_FONT, align=LEFT)
    ws.merge_cells("F14:H14")
    ws.row_dimensions[14].height = 40
    dv_sop = DataValidation(type="list", formula1='"已收集,部分收集,无SOP可交,待收集"', allow_blank=True)
    ws.add_data_validation(dv_sop)
    dv_sop.add("C14")

    merge_write(
        ws, 15, 1, 2,
        "收集说明 / 缺口备注\n（缺哪些人、哪些文档）",
        fill=SECTION_FILL, font=LABEL_FONT, align=LEFT_TOP, height=56,
    )
    merge_write(ws, 15, 3, 8, "", fill=INPUT_FILL, font=CELL_FONT, align=LEFT_TOP, height=56)

    merge_write(
        ws, 16, 1, 8,
        "【提交说明】本部门同事在各自《个人工作事项填写表》每条事项填写「对应标准作业文档名」。请您按名单统一收集全部 SOP/作业指导书，放入同一个文件夹后打成一个压缩包，与本表一并提交。建议命名：标准作业SOP-公司名-部门名.zip。无人有 SOP 则选「无SOP可交」，不必交空包。",
        fill=SECTION_FILL, font=HINT_FONT, align=LEFT_TOP, height=56,
    )

    merge_write(ws, 18, 1, 8, "备注", fill=SECTION_FILL, font=LABEL_FONT, align=LEFT, height=26)
    merge_write(ws, 19, 1, 8, "", fill=INPUT_FILL, font=CELL_FONT, align=LEFT_TOP, height=48)
    setup_sheet(ws, last_col="H", freeze="A3")

    # ---- 02岗位编制 ----
    ws = wb.create_sheet("02岗位编制")
    headers = ["序号", "岗位名称", "人数", "主要职责（一句话）", "上班时段/班次", "直接上级\n（可写岗位）", "是否直接\n处理业务", "备注"]
    set_widths(ws, [6, 18, 8, 42, 18, 16, 12, 22])
    merge_write(ws, 1, 1, 8, "组织信息填写表 · 岗位编制", fill=SECTION_FILL, font=TITLE_FONT, align=CENTER, height=30)
    merge_write(
        ws, 2, 1, 8,
        "列出本部门岗位即可；人数可约数。「是否直接处理业务」= 是否日常操作单据/系统/往来业务（非纯管理协调），下拉：是 / 否",
        font=DOC_FONT, align=LEFT, height=28,
    )
    write_header_row(ws, 3, headers)
    blank_input_rows(ws, 4, 20, 8, height=38)
    for i, r in enumerate(range(4, 21), 1):
        ws.cell(r, 1, i).alignment = CENTER
    dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("G4:G20")
    setup_sheet(ws, last_col="H")

    # ---- 03业务量基线：短表头 + 说明行 ----
    ws = wb.create_sheet("03业务量基线")
    headers = [
        "序号",
        "业务事项/产物",
        "计量单位",
        "日均量",
        "周均量",
        "峰值量及何时",
        "单次耗时\n(分钟)",
        "参与\n人数",
        "差错返工\n约占比%",
        "数字来源",
        "备注",
    ]
    # 加宽事项与备注，数值列适中
    set_widths(ws, [5, 34, 12, 10, 10, 18, 11, 8, 11, 14, 22])
    merge_write(ws, 1, 1, 11, "组织信息填写表 · 业务量基线", fill=SECTION_FILL, font=TITLE_FONT, align=CENTER, height=30)
    merge_write(
        ws, 2, 1, 11,
        "按「本部门整体」填工作量最大的几项（不是个人视角；个人明细在《个人工作事项填写表》）。计量单位手写（票/单/件等）。数字可估算，来源写「台账/系统导出/估算」。峰值写「量+时段」（如月末）。本表不设下拉。",
        font=DOC_FONT, align=LEFT, height=36,
    )
    write_header_row(ws, 3, headers, height=52)
    blank_input_rows(ws, 4, 16, 11, height=38)
    for i, r in enumerate(range(4, 17), 1):
        ws.cell(r, 1, i).alignment = CENTER
    setup_sheet(ws, last_col="K")

    # ---- 04系统与产物 ----
    ws = wb.create_sheet("04系统与产物清单")
    headers = [
        "序号",
        "系统/工具名称",
        "类型",
        "主要用途",
        "使用岗位",
        "能否\n导出",
        "有无系统\n接口(API)",
        "关联产物\n(单据/报表/台账)",
        "样例\n可否提供",
        "近期\n是否改版",
        "备注",
    ]
    set_widths(ws, [5, 20, 14, 26, 14, 9, 9, 24, 12, 10, 18])
    merge_write(ws, 1, 1, 11, "组织信息填写表 · 系统与产物清单", fill=SECTION_FILL, font=TITLE_FONT, align=CENTER, height=30)
    merge_write(
        ws, 2, 1, 11,
        "列出本部门日常用到的系统/软件/表格工具。类型、用途、岗位、关联产物手写。不知「有无接口」可选「不知」。勿填账号密码。",
        font=DOC_FONT, align=LEFT, height=28,
    )
    write_header_row(ws, 3, headers, height=52)
    blank_input_rows(ws, 4, 20, 11, height=38)
    for i, r in enumerate(range(4, 21), 1):
        ws.cell(r, 1, i).alignment = CENTER
    # 仅保留可规范的是非类下拉；类型/用途等自由填写
    for col, opts in [
        ("F", '"能,部分能,不能,不知"'),
        ("G", '"有,无,不知"'),
        ("I", '"可提供,仅现场看,暂不提供"'),
        ("J", '"是,否,计划中,不知"'),
    ]:
        dvx = DataValidation(type="list", formula1=opts, allow_blank=True)
        ws.add_data_validation(dvx)
        dvx.add(f"{col}4:{col}20")
    setup_sheet(ws, last_col="K")

    move_sheet_to_end(wb)
    wb.active = 0
    out = OUT_DIR / ORG_NAME
    wb.save(out)
    return out


def build_personal_workbook():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "填写说明"
    set_widths(ws, [12, 14, 72, 12, 12])
    merge_write(ws, 1, 1, 5, "个人工作事项填写表 · 填写说明", fill=SECTION_FILL, font=TITLE_FONT, align=CENTER, height=34)
    merge_write(
        ws, 2, 1, 5,
        "请填写你本人日常工作；与《组织信息填写表》成套。白色为待填。",
        font=DOC_FONT, align=CENTER, height=22,
    )
    rows = [
        ("用途", "写清你日常做什么、大概多久一次、哪里最费时/最烦，方便进场前先摸清情况。"),
        ("谁填", "本部门日常做事的同事；每人填自己的。"),
        ("成套关系", "部门负责人填《组织信息填写表》；你填本表。一份部门表 + 若干个人表。"),
        ("怎么填", "一件事一行；同一人可多行。写具体动作，不要只写岗位名。"),
        ("脱敏", "勿贴业务往来方完整姓名、未脱敏单号、账号密码；可用「往来方A」「单号***」。"),
        ("耗时频率", "频率请写清次数，如「每天 3 次」「每周 2 次」；单次耗时按典型情况估算，高峰可写备注。"),
        ("SOP 文档名", "在「02工作事项明细」每一行事项旁填对应标准作业文档名（没有写「无」）。纸质/电子文件交给部门负责人统一打包。"),
        ("痛点速填", "「03痛点速填」必填：写出最耗时/最烦的三件事及原因，不可跳过。"),
    ]
    r = 4
    for title, body in rows:
        merge_write(ws, r, 1, 2, title, fill=SECTION_FILL, font=LABEL_FONT, align=CENTER, height=42)
        merge_write(ws, r, 3, 5, body, fill=INPUT_FILL, font=CELL_FONT, align=LEFT_TOP, height=42)
        r += 1
    merge_write(
        ws, r + 1, 1, 5,
        "交回本表 Excel 即可（建议名：个人工作事项填写表-公司名-部门名-姓名.xlsx）。SOP 压缩包由部门负责人统一提交。",
        font=HINT_FONT, align=LEFT_TOP, height=36,
    )
    setup_sheet(ws, last_col="E", landscape=False, freeze=None)

    ws = wb.create_sheet("01填写人信息")
    set_widths(ws, [20, 22, 16, 22, 18, 22, 14, 20])
    merge_write(ws, 1, 1, 8, "个人工作事项填写表 · 填写人信息", fill=SECTION_FILL, font=TITLE_FONT, align=CENTER, height=30)
    merge_write(ws, 2, 1, 8, "请本人填写；与《组织信息填写表》成套。多人合用一文件时在明细用姓名区分，或每人复制本文件", font=DOC_FONT, align=CENTER, height=22)

    label_value_row(ws, 3, [("公司名称", 1, 2), ("本部门", 3, 4), ("填表日期", 5, 6), ("姓名", 7, 8)], height=32)
    label_value_row(ws, 4, [("岗位", 1, 2), ("入职时长(约)", 3, 4), ("手机/微信(可选)", 5, 6), ("直属上级岗位", 7, 8)], height=32)
    label_value_row(
        ws, 5,
        [("主要服务对象\n(内部同事/外部往来方)", 1, 2, 4), ("日常最忙的时段", 5, 6, 8)],
        height=36,
    )

    merge_write(ws, 7, 1, 8, "一句话：你最希望先少做、少烦的一件事是什么？", fill=HEADER_FILL, font=HEADER_FONT, align=LEFT, height=28)
    merge_write(ws, 8, 1, 8, "", fill=INPUT_FILL, font=CELL_FONT, align=LEFT_TOP, height=56)

    paint(ws, 10, 10, 1, 4, fill=SECTION_FILL, font=LABEL_FONT, align=LEFT)
    ws.cell(10, 1, "若后续有改进试点，是否愿意先一起试")
    ws.merge_cells("A10:D10")
    paint(ws, 10, 10, 5, 8, fill=INPUT_FILL, font=CELL_FONT, align=LEFT)
    ws.merge_cells("E10:H10")
    ws.row_dimensions[10].height = 32
    dv = DataValidation(type="list", formula1='"是,否,看情况"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("E10")
    setup_sheet(ws, last_col="H", freeze="A3")

    # 工作事项明细：每事项一行；SOP 文档名跟事项走
    ws = wb.create_sheet("02工作事项明细")
    headers = [
        "序号",
        "姓名",
        "工作事项\n（具体动作）",
        "触发条件",
        "完成标准",
        "频率\n(如每天3次)",
        "单次耗时\n(分钟)",
        "系统/工具",
        "规则/判断\n(可写比例)",
        "是否\n重复录入",
        "最常见出错/异常",
        "异常\n谁处理",
        "痛点\n程度",
        "对应标准作业\n文档名(无则写无)",
        "备注",
    ]
    set_widths(ws, [5, 10, 26, 16, 16, 12, 10, 14, 12, 9, 16, 10, 8, 22, 14])
    merge_write(ws, 1, 1, 15, "个人工作事项填写表 · 工作事项明细", fill=SECTION_FILL, font=TITLE_FONT, align=CENTER, height=30)
    merge_write(
        ws, 2, 1, 15,
        "一人可多行。每条事项单独写「对应标准作业文档名」（作业指导书/SOP 文件名）；没有写「无」。频率等手写。仅「是否重复录入」「痛点程度」有下拉。"
        " 实体 SOP 文件交部门负责人统一收集打包，个人只需交本表。",
        font=DOC_FONT, align=LEFT, height=42,
    )
    write_header_row(ws, 3, headers, height=54)
    blank_input_rows(ws, 4, 25, 15, height=42)
    for i, r in enumerate(range(4, 26), 1):
        ws.cell(r, 1, i).alignment = CENTER

    # 列：J=是否重复录入(10), M=痛点程度(13)
    for col, opts in [
        ("J", '"是,否,部分"'),
        ("M", '"高,中,低"'),
    ]:
        dvx = DataValidation(type="list", formula1=opts, allow_blank=True)
        ws.add_data_validation(dvx)
        dvx.add(f"{col}4:{col}25")
    setup_sheet(ws, last_col="O")

    ws = wb.create_sheet("03痛点速填")
    set_widths(ws, [8, 42, 48, 16])
    merge_write(ws, 1, 1, 4, "个人工作事项填写表 · 痛点速填（必填）", fill=SECTION_FILL, font=TITLE_FONT, align=CENTER, height=30)
    merge_write(ws, 2, 1, 4, "必填。请写出日常最耗时 / 最烦的三件事及原因；可与「02工作事项明细」对照，但本页不可留空。", font=DOC_FONT, align=CENTER, height=28)
    write_header_row(ws, 3, ["排名", "最耗时 / 最烦的事", "为什么烦（重复/易错/等人/跨系统…）", "姓名"], height=36)
    for i, r in enumerate(range(4, 7), 1):
        for c in range(1, 5):
            cell = ws.cell(r, c, i if c == 1 else "")
            cell.fill = INPUT_FILL
            cell.font = CELL_FONT
            cell.alignment = CENTER if c == 1 else LEFT_TOP
            cell.border = THIN
        ws.row_dimensions[r].height = 48
    merge_write(ws, 8, 1, 4, "其他想补充的话", fill=SECTION_FILL, font=LABEL_FONT, align=LEFT, height=26)
    merge_write(ws, 9, 1, 4, "", fill=INPUT_FILL, font=CELL_FONT, align=LEFT_TOP, height=64)
    setup_sheet(ws, last_col="D", landscape=False, freeze="A4")

    move_sheet_to_end(wb)
    wb.active = 0
    out = OUT_DIR / PERSONAL_NAME
    wb.save(out)
    return out


def main():
    a = build_org_workbook()
    b = build_personal_workbook()
    print(f"OK: {a}")
    print(f"OK: {b}")


if __name__ == "__main__":
    main()
