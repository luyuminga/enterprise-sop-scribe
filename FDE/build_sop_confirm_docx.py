#!/usr/bin/env python3
"""Export SOP confirmation Word from flow MD (01 mermaid + 02 SOP)."""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("缺少 python-docx：pip install python-docx", file=sys.stderr)
    sys.exit(1)


CONFIRM_HINT_TOP = (
    "请先看下方【业务流程图】：主干与分支是否与您实际做法一致？"
    "图核对完后再看业务背景与步骤；有出入请直接修改或回复修改点。"
)
CONFIRM_HINT_FOOT = (
    "请再次确认：流程图 + 下列说明是否与您实际做法一致？"
    "确认通过后本流程方算定稿，方可派生知识库与自动化分析。"
)


def extract_meta(sop_text: str, flow_dir: Path) -> dict[str, str]:
    meta = {
        "事项名称": flow_dir.name.split("-", 1)[-1] if "-" in flow_dir.name else flow_dir.name,
        "适用岗位": "",
        "确认状态": "",
    }
    for key in ("事项名称", "适用岗位", "确认状态"):
        m = re.search(rf"\|\s*{re.escape(key)}\s*\|\s*([^|]+)\|", sop_text)
        if m:
            val = m.group(1).strip()
            if val and not val.startswith("{{"):
                meta[key] = val
    return meta


def set_run_font(run, *, east_asia: str, size_pt: float, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = "Times New Roman"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), east_asia)


def add_para(doc: Document, text: str, *, bold: bool = False, size: float = 12) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, east_asia="宋体", size_pt=size, bold=bold)
    p.paragraph_format.line_spacing = 1.5
    return p


def extract_mermaid_blocks(md: str) -> list[str]:
    return re.findall(r"```mermaid\s*([\s\S]*?)```", md, flags=re.IGNORECASE)


def render_mermaid(mmd: str, out_png: Path) -> bool:
    mmdc = shutil.which("mmdc")
    if not mmdc:
        return False
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "flow.mmd"
        src.write_text(mmd.strip() + "\n", encoding="utf-8")
        try:
            subprocess.run(
                [mmdc, "-i", str(src), "-o", str(out_png), "-b", "white"],
                check=True,
                capture_output=True,
                text=True,
            )
            return out_png.is_file()
        except (subprocess.CalledProcessError, FileNotFoundError) as e:
            print(f"mmdc 渲染失败: {e}", file=sys.stderr)
            return False


def md_to_docx_body(doc: Document, md: str) -> None:
    """Lightweight MD → Word: headings, tables, paragraphs; skip mermaid fences."""
    lines = md.splitlines()
    i = 0
    in_fence = False
    fence_lang = ""
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            lang = line.strip()[3:].strip().lower()
            if not in_fence:
                in_fence = True
                fence_lang = lang
                i += 1
                continue
            in_fence = False
            fence_lang = ""
            i += 1
            continue
        if in_fence:
            if fence_lang == "mermaid":
                i += 1
                continue
            add_para(doc, line, size=10.5)
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^[\s\-:|]+$", lines[i].strip().strip("|").replace("|", "")):
                    rows.append(row)
                i += 1
                if i < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i].strip()):
                    i += 1
                    continue
                if i >= len(lines) or not lines[i].strip().startswith("|"):
                    break
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(cols):
                        cell = table.cell(ri, ci)
                        cell.text = row[ci] if ci < len(row) else ""
                        for p in cell.paragraphs:
                            for run in p.runs:
                                set_run_font(run, east_asia="宋体", size_pt=10.5, bold=(ri == 0))
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            p = doc.add_paragraph()
            run = p.add_run(text)
            size = {1: 16, 2: 14, 3: 12}[level]
            set_run_font(run, east_asia="黑体", size_pt=size, bold=True)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        add_para(doc, line.strip())
        i += 1


def add_heading(doc: Document, text: str, size: float = 14) -> None:
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), east_asia="黑体", size_pt=size, bold=True)


def build(flow_dir: Path, out: Path) -> int:
    sop = flow_dir / "02-SOP作业指导书.md"
    flow_md = flow_dir / "01-业务流程图.md"
    if not sop.is_file():
        print(f"缺少文件: {sop}", file=sys.stderr)
        return 2
    if not flow_md.is_file():
        print(f"警告: 缺少 {flow_md.name}，流程图节将为空占位。", file=sys.stderr)

    sop_text = sop.read_text(encoding="utf-8")
    flow_text = flow_md.read_text(encoding="utf-8") if flow_md.is_file() else ""
    meta = extract_meta(sop_text, flow_dir)

    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)

    # --- 标识 ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("业务流程确认稿")
    set_run_font(tr, east_asia="黑体", size_pt=18, bold=True)

    ident = doc.add_paragraph()
    ident.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ir = ident.add_run(
        f"事项：{meta['事项名称']}"
        + (f"　｜　岗位：{meta['适用岗位']}" if meta["适用岗位"] else "")
        + (f"　｜　状态：{meta['确认状态']}" if meta["确认状态"] else "")
    )
    set_run_font(ir, east_asia="宋体", size_pt=12, bold=True)

    hint = doc.add_paragraph()
    hr = hint.add_run(CONFIRM_HINT_TOP)
    set_run_font(hr, east_asia="宋体", size_pt=12, bold=True)
    hr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # --- 一、业务流程图（第一要义：直观确认）---
    add_heading(doc, "一、业务流程图（请先核对）", 16)
    add_para(
        doc,
        "本图是确认重点：路径、判断分支、异常回路是否与您平时做的一致。"
        "FDE 亦可据此快速建立对整条业务流的整体认识。",
        size=10.5,
    )

    blocks = extract_mermaid_blocks(flow_text)
    if not blocks:
        add_para(doc, "（未在 01-业务流程图.md 中找到 mermaid 代码块；请补全后再导出。）")
    else:
        with tempfile.TemporaryDirectory() as td:
            png = Path(td) / "flow.png"
            ok = render_mermaid(blocks[0], png)
            if ok:
                doc.add_picture(str(png), width=Cm(16.0))
            else:
                print("警告: 未找到 mmdc 或渲染失败，Word 中流程图降级为文字。", file=sys.stderr)
                add_para(
                    doc,
                    "（流程图出图失败：请安装 @mermaid-js/mermaid-cli 后重导。"
                    "以下为 Mermaid 原文，供临时核对。）",
                )
                add_para(doc, blocks[0][:3000], size=9)
            # 节点说明（mermaid 外的正文）
            rest = re.sub(r"```mermaid\s*[\s\S]*?```", "", flow_text, count=1, flags=re.I).strip()
            rest = re.sub(r"^#.*$", "", rest, flags=re.M).strip()
            if rest:
                add_heading(doc, "图注 / 节点说明", 12)
                md_to_docx_body(doc, rest)

    # --- 二及以后：业务背景与步骤（图之后才展开）---
    add_heading(doc, "二、业务背景与作业说明", 16)
    add_para(
        doc,
        "以下为背景（目的/SIPOC/触发等）与可执行步骤，供图核对后的细节确认，并便于 FDE 落地理解。",
        size=10.5,
    )
    md_to_docx_body(doc, sop_text)

    foot = doc.add_paragraph()
    fr = foot.add_run(CONFIRM_HINT_FOOT)
    set_run_font(fr, east_asia="宋体", size_pt=12, bold=True)
    fr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"已生成: {out}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="从业务流 MD 导出 Word 确认稿")
    ap.add_argument("--flow-dir", required=True, help="flows/0N-短名 目录")
    ap.add_argument("--out", default="", help="输出 docx 路径")
    args = ap.parse_args()
    flow_dir = Path(args.flow_dir).expanduser().resolve()
    out = (
        Path(args.out).expanduser().resolve()
        if args.out
        else flow_dir / "02-SOP作业指导书-确认稿.docx"
    )
    return build(flow_dir, out)


if __name__ == "__main__":
    sys.exit(main())
